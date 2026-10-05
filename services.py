"""Glue between the UI and the Sarvam pipeline, with caching and an offline fallback.

Every function takes `live`. live=True calls Sarvam (needs SARVAM_API_KEY in .env);
live=False only uses caches, bundled demo data and built-in text, so a demo never depends on wifi.
"""
import base64
import hashlib
import json
import os
import pathlib

from claims import Claim
from normalize import normalize_all

ROOT = pathlib.Path(__file__).resolve().parent
WORK = ROOT / "work"
DEMO_DIR = ROOT / "data" / "demo"
SAMPLE_AUDIO_DIR = ROOT / "hacksprintdata"
SAMPLE_XLSX = SAMPLE_AUDIO_DIR / "voice_eval_multilingual.xlsx"
LANG_CODES = {"en": "en-IN", "hi": "hi-IN", "ta": "ta-IN"}
SPEAKER = os.environ.get("BULBUL_SPEAKER", "shubh")

_client = None


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def api_key():
    try:
        from doc_engine.config import get_api_key
        return get_api_key()
    except Exception:
        return None


def live_available() -> bool:
    try:
        import sarvamai  # noqa: F401
    except ImportError:
        return False
    return api_key() is not None


def client():
    global _client
    if _client is None:
        from sarvamai import SarvamAI
        _client = SarvamAI(api_subscription_key=api_key())
        os.environ.setdefault("SARVAM_API_KEY", api_key())
    return _client


def _jload(p: pathlib.Path):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def _jsave(p: pathlib.Path, data):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------- agreements (P1) ----------------

def demo_agreements() -> dict:
    """name -> facts dict, from bundled fixtures (ground-truth values of the synthetic agreements)."""
    out = {}
    for p in sorted(DEMO_DIR.glob("*.facts.json")):
        out[p.name.removesuffix(".facts.json")] = json.loads(p.read_text(encoding="utf-8"))
    return out


def facts_for_pdf(data: bytes, name: str, live: bool):
    """-> (facts, source) or (None, reason). source: demo | cache | live."""
    aid = _sha(data)[:12]
    for facts in demo_agreements().values():
        if facts["agreement_id"] == aid:
            return facts, "demo"
    from doc_engine import cache
    cached = cache.load(aid, "facts")
    if cached:
        return cached, "cache"
    if not live:
        return None, "needs_live"
    from doc_engine.build import build_agreement_facts
    path = WORK / "uploads" / f"{aid}.pdf"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return build_agreement_facts(str(path)), "live"


# ---------------- voice (P2) ----------------

def sample_clips() -> list:
    """Clips from the answer-key workbook: id, language, spoken script, expected claims, audio path."""
    if not SAMPLE_XLSX.exists():
        return []
    import pipeline
    lang, expected = pipeline.load_key(str(SAMPLE_XLSX))
    import openpyxl
    wb = openpyxl.load_workbook(SAMPLE_XLSX)
    out = []
    for row in wb["Clips to Record"].iter_rows(min_row=2, max_col=3, values_only=True):
        cid, lg, script = row
        if cid is None or lg is None or cid not in expected:
            continue
        audio = next((p for p in SAMPLE_AUDIO_DIR.glob(f"{str(cid).lower()}.*") if p.suffix != ".xlsx"), None)
        out.append(dict(id=cid, language=lg, script=script, expected=expected[cid], audio=audio))
    return out


def claims_from_audio(data: bytes, filename: str, live: bool, lang_hint=None):
    """Audio bytes -> dict(transcript, claims, source). Cached by audio hash.
    Returns None if not cached and live is off."""
    sha = _sha(data)[:16]
    cache = WORK / "voice" / f"{sha}.json"
    hit = _jload(cache)
    if hit:
        return dict(transcript=hit["transcript"], claims=normalize_all(hit["raw"]), source="cache")
    if not live:
        return None
    import pipeline
    cl = client()
    path = WORK / "uploads" / f"{sha}{pathlib.Path(filename).suffix or '.wav'}"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    code = pipeline.LANG_CODE.get(lang_hint, "unknown")
    try:
        transcript = pipeline.transcribe(cl, path, code, "codemix")
    except Exception:
        transcript = pipeline.transcribe(cl, path, "en-IN", "codemix")
    raw = pipeline.parse_claims(cl, transcript, sha)
    _jsave(cache, dict(transcript=transcript, raw=raw))
    return dict(transcript=transcript, claims=normalize_all(raw), source="live")


def claims_from_sample(clip: dict, live: bool):
    """Sample clip -> same shape as claims_from_audio.
    live + audio file: run the real pipeline on the recording. Otherwise use the written
    script and its answer-key claims (clearly reported as source='sample')."""
    if live and clip["audio"]:
        res = claims_from_audio(clip["audio"].read_bytes(), clip["audio"].name, True, clip["language"])
        if res:
            return res
    raw = [dict(term=e["term"], value=e["value"], unit=e["unit"] or "none", hedged=e["hedged"], quote=clip["script"])
           for e in clip["expected"]]
    return dict(transcript=clip["script"], claims=normalize_all(raw), source="sample")


# ---------------- Translate (templates) and Bulbul (audio) ----------------

def translate_template(text: str, lang: str, live: bool, placeholders=("{term}", "{said}", "{actual}")):
    """Sarvam Translate for a sentence template. Returns the translation, or None
    (caller falls back to built-in text). Placeholders must survive translation."""
    if lang == "en":
        return text
    cache = WORK / "translate" / f"{_sha((text + lang).encode())[:16]}.json"
    hit = _jload(cache)
    if hit:
        return hit["text"]
    if not live:
        return None
    try:
        r = client().text.translate(input=text, source_language_code="en-IN", target_language_code=LANG_CODES[lang],
                                    model="sarvam-translate:v1")
        out = getattr(r, "translated_text", None)
    except Exception:
        return None
    if not out or any((p in text) != (p in out) for p in placeholders):
        return None
    _jsave(cache, dict(src=text, lang=lang, text=out))
    return out


def speak(text: str, lang: str, live: bool):
    """Bulbul v3 -> WAV bytes (cached), or None if unavailable."""
    cache = WORK / "audio" / f"{_sha((text + lang + SPEAKER).encode())[:16]}.wav"
    if cache.exists():
        return cache.read_bytes()
    if not live:
        return None
    try:
        r = client().text_to_speech.convert(text=text, target_language_code=LANG_CODES[lang],
                                            model="bulbul:v3", speaker=SPEAKER)
        wav = base64.b64decode("".join(r.audios))
    except Exception:
        return None
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_bytes(wav)
    return wav
