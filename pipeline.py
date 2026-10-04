"""P2 pipeline: audio -> Saaras transcript -> sarvam-105b parser -> normalize -> score.

Setup:   pip install -U sarvamai openpyxl pydantic ; export SARVAM_API_KEY=...
Layout:  audio/clip01_priya.wav, audio/clipT1_x.wav, ...  (clip id at the start of the name)

Step 1:  python pipeline.py transcribe        # runs transcribe + codemix, prints them side by side
Step 2:  python pipeline.py run --mode codemix   # parse + normalize + score; writes eval_results.csv

Calls follow Sarvam's quickstart (saaras:v4 STT, sarvam-105b chat). Everything is cached
under work/, so re-running costs nothing for files already done.
"""
import argparse
import csv
import json
import os
import pathlib
import re
from collections import defaultdict

from normalize import normalize_claim

LANG_CODE = {"English": "en-IN", "Hinglish": "hi-IN", "Tamil": "ta-IN"}
WORK = pathlib.Path("work")

PARSER_PROMPT = """You extract rental-agreement claims from a spoken transcript (English, Hinglish or Tamil).
Return ONLY JSON: {"claims": [{"term": ..., "value": ..., "unit": ..., "hedged": ..., "quote": ...}]}

term is one of: monthly_rent, security_deposit, lock_in_period, notice_period, maintenance_charges, deposit_refund_period
unit is one of: inr, days, months, months_of_rent, none
value is the number as a JSON number (convert spoken number words yourself); null if the speaker gives no number.
hedged is true if the speaker is unsure (around, maybe, I think, approximately, irukkalaam, shayad).
quote is the exact words from the transcript that support the claim.
Rules:
- "N months' deposit" -> unit months_of_rent; a rupee amount for the deposit -> unit inr.
- "no lock-in" -> lock_in_period, value 0, unit months.
- If maintenance is said to be included in the rent, value null, unit none.
- One claim per term mentioned. Do not invent terms that were not said.

Examples (different clips from the test set):
Transcript: The rent is twenty-five thousand and the notice is one month.
{"claims": [{"term": "monthly_rent", "value": 25000, "unit": "inr", "hedged": false, "quote": "The rent is twenty-five thousand"}, {"term": "notice_period", "value": 1, "unit": "months", "hedged": false, "quote": "the notice is one month"}]}
Transcript: Deposit teen mahine ka hai, shayad.
{"claims": [{"term": "security_deposit", "value": 3, "unit": "months_of_rent", "hedged": true, "quote": "Deposit teen mahine ka hai, shayad"}]}
"""


# ---------- answer key ----------

def clip_id_from_name(name: str):
    m = re.match(r"(?:clip[_-]?)?(T\d+|\d+)", name, re.I)
    if not m:
        return None
    s = m.group(1).upper()
    return s if s.startswith("T") else int(s)


def load_key(xlsx_path: str):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    lang = {}
    for row in wb["Clips to Record"].iter_rows(min_row=2, max_col=2, values_only=True):
        if row[0] is not None and row[1] in LANG_CODE:
            lang[row[0]] = row[1]
    expected = defaultdict(list)
    for row in wb["Answer Key"].iter_rows(min_row=2, max_col=6, values_only=True):
        cid, _lang, term, value, unit, hedged = row
        if cid is None or term is None:
            continue
        expected[cid].append(dict(term=term, value=value, unit=unit, hedged=(hedged == "=TRUE()")))
    return lang, expected


# ---------- Sarvam calls ----------

def make_client():
    from sarvamai import SarvamAI
    return SarvamAI(api_subscription_key=os.environ["SARVAM_API_KEY"])


def transcribe(client, path: pathlib.Path, lang_code: str, mode: str) -> str:
    out = WORK / "transcripts" / f"{path.stem}.{mode}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        return json.loads(out.read_text())["transcript"]
    with open(path, "rb") as f:
        r = client.speech_to_text.transcribe(file=f, model="saaras:v4", mode=mode, language_code=lang_code)
    out.write_text(json.dumps({"file": path.name, "mode": mode, "transcript": r.transcript}, ensure_ascii=False, indent=2))
    return r.transcript


def extract_json(text: str):
    """Pull the first JSON object out of a model reply (strips <think> blocks and code fences)."""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    text = re.sub(r"```(?:json)?", "", text)
    start = text.find("{")
    if start == -1:
        return None
    depth = 0
    for i in range(start, len(text)):
        depth += text[i] == "{"
        depth -= text[i] == "}"
        if depth == 0:
            try:
                return json.loads(text[start:i + 1])
            except json.JSONDecodeError:
                return None
    return None


def parse_claims(client, transcript: str, cache_name: str, attempts: int = 2) -> list[dict]:
    out = WORK / "parsed" / f"{cache_name}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        return json.loads(out.read_text())
    claims = None
    for _ in range(attempts):
        r = client.chat.completions(
            model="sarvam-105b",
            messages=[{"role": "system", "content": PARSER_PROMPT},
                      {"role": "user", "content": f"Transcript: {transcript}"}],
        )
        data = extract_json(r.choices[0].message.content or "")
        if isinstance(data, dict) and isinstance(data.get("claims"), list):
            claims = data["claims"]
            break
    claims = claims or []
    out.write_text(json.dumps(claims, ensure_ascii=False, indent=2))
    return claims


# ---------- scoring ----------

def score(expected: list[dict], got_claims) -> list[dict]:
    """One row per expected claim: term_ok = term found, value_ok = value AND unit match."""
    rows = []
    for e in expected:
        g = next((c for c in got_claims if c.term == e["term"]), None)
        term_ok = g is not None
        value_ok = bool(g) and g.unit == e["unit"] and (
            (g.value is None and e["value"] is None) or
            (g.value is not None and e["value"] is not None and abs(g.value - float(e["value"])) < 1e-9))
        rows.append(dict(term=e["term"], exp_value=e["value"], exp_unit=e["unit"],
                         got_value=g.value if g else None, got_unit=g.unit if g else None,
                         term_ok="Y" if term_ok else "N", value_ok="Y" if value_ok else "N",
                         issues=";".join(g.issues) if g else "term_missing"))
    return rows


# ---------- commands ----------

def cmd_transcribe(args):
    client = make_client()
    lang, _ = load_key(args.key)
    for f in sorted(pathlib.Path(args.audio).glob("*")):
        cid = clip_id_from_name(f.name)
        language = lang.get(cid)
        if not language:
            print(f"skip {f.name}: clip id {cid!r} not in sheet")
            continue
        print(f"\n== {f.name} ({language})")
        for mode in ("transcribe", "codemix"):
            print(f"  {mode:10s} {transcribe(client, f, LANG_CODE[language], mode)}")


def cmd_run(args):
    client = make_client()
    lang, expected = load_key(args.key)
    results = []
    for f in sorted(pathlib.Path(args.audio).glob("*")):
        cid = clip_id_from_name(f.name)
        language = lang.get(cid)
        if not language or cid not in expected:
            print(f"skip {f.name}")
            continue
        text = transcribe(client, f, LANG_CODE[language], args.mode)
        raw = parse_claims(client, text, f"{f.stem}.{args.mode}")
        got = [normalize_claim(c) for c in raw if isinstance(c, dict)]
        for row in score(expected[cid], got):
            results.append(dict(clip_id=cid, file=f.name, language=language, **row, transcript=text))

    if not results:
        print("nothing scored")
        return
    with open("eval_results.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(results[0]))
        w.writeheader()
        w.writerows(results)

    by_lang = defaultdict(lambda: [0, 0])
    for r in results:
        by_lang[r["language"]][0] += r["term_ok"] == "Y" and r["value_ok"] == "Y"
        by_lang[r["language"]][1] += 1
    print("\nAccuracy (term AND value correct):")
    for language, (ok, n) in by_lang.items():
        print(f"  {language:9s} {ok}/{n} = {ok / n:.0%}")
    print("Failures:")
    for r in results:
        if r["value_ok"] != "Y":
            print(f"  {r['file']}: {r['term']} expected {r['exp_value']} {r['exp_unit']}, "
                  f"got {r['got_value']} {r['got_unit']} {r['issues']}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("command", choices=["transcribe", "run"])
    ap.add_argument("--audio", default="audio")
    ap.add_argument("--key", default="voice_eval_multilingual.xlsx")
    ap.add_argument("--mode", default="codemix", choices=["transcribe", "codemix", "verbatim", "translit"])
    a = ap.parse_args()
    {"transcribe": cmd_transcribe, "run": cmd_run}[a.command](a)
